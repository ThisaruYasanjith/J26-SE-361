
from pathlib import Path
from collections import Counter
import hashlib
import json

from openpyxl import load_workbook


# Resolve all paths relative to the Code Cleanliness folder.
COMPONENT_ROOT = Path(__file__).resolve().parents[2]

DATASET_ROOT = (
    COMPONENT_ROOT
    / "datasets"
    / "raw"
    / "LLM4Refactoring"
    / "data"
    / "dataset"
)

METADATA_FILE = DATASET_ROOT / "ref-Dataset.xlsx"
OUTPUT_DIR = COMPONENT_ROOT / "datasets" / "processed"
OUTPUT_FILE = OUTPUT_DIR / "llm4refactoring_inventory.jsonl"


def normalize_id(value):
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def sha256_file(path):
    digest = hashlib.sha256()

    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(65536), b""):
            digest.update(chunk)

    return digest.hexdigest()


def build_inventory():
    if not METADATA_FILE.is_file():
        raise FileNotFoundError(
            f"Metadata spreadsheet not found: {METADATA_FILE}"
        )

    workbook = load_workbook(
        METADATA_FILE,
        read_only=True,
        data_only=True
    )

    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)

        required = [
            "Refactoring Type",
            "ID",
            "Project Name",
            "Commit SHA",
            "File",
            "Refactoring",
        ]

        missing = [h for h in required if h not in headers]

        if missing:
            raise ValueError(
                f"Missing metadata columns: {missing}"
            )

        indexes = {
            name: headers.index(name)
            for name in required
        }

        records = []
        seen_keys = set()
        current_category = None

        for spreadsheet_row, values in enumerate(rows, start=2):

            if not any(value is not None for value in values):
                continue

            def cell(name):
                return values[indexes[name]]

            raw_category = cell("Refactoring Type")

            if raw_category is not None and str(raw_category).strip():
                current_category = str(raw_category).strip()

            category = current_category or ""

            example_id = normalize_id(cell("ID"))
            key = (category, example_id)

            if not category or not example_id:
                raise ValueError(
                    f"Missing category or ID on row {spreadsheet_row}"
                )

            if key in seen_keys:
                raise ValueError(
                    f"Duplicate category and ID: {key}"
                )

            seen_keys.add(key)

            example_folder = (
                DATASET_ROOT / category / example_id
            )

            before = example_folder / "before.java"
            after = example_folder / "after.java"

            before_exists = before.is_file()
            after_exists = after.is_file()

            before_size = (
                before.stat().st_size if before_exists else 0
            )

            after_size = (
                after.stat().st_size if after_exists else 0
            )

            status = (
                "valid_pair"
                if before_size > 0 and after_size > 0
                else "missing_or_empty"
            )

            record = {
                "source": "LLM4Refactoring",
                "spreadsheet_row": spreadsheet_row,
                "refactoring_type": category,
                "example_id": example_id,
                "project_name": cell("Project Name"),
                "commit_sha": cell("Commit SHA"),
                "original_file": cell("File"),
                "refactoring_metadata": cell("Refactoring"),
                "extra_column_value": (
                    values[6] if len(values) > 6 else None
                ),
                "before_path": str(
                    before.relative_to(COMPONENT_ROOT)
                ),
                "after_path": str(
                    after.relative_to(COMPONENT_ROOT)
                ),
                "before_bytes": before_size,
                "after_bytes": after_size,
                "before_sha256": (
                    sha256_file(before) if before_exists else None
                ),
                "after_sha256": (
                    sha256_file(after) if after_exists else None
                ),
                "status": status,
            }

            records.append(record)

    finally:
        workbook.close()

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True
    )

    with OUTPUT_FILE.open(
        "w",
        encoding="utf-8"
    ) as output:

        for record in records:
            output.write(
                json.dumps(
                    record,
                    ensure_ascii=False,
                    default=str
                ) + "\n"
            )

    categories = Counter(
        record["refactoring_type"]
        for record in records
    )

    statuses = Counter(
        record["status"]
        for record in records
    )

    print("\nLLM4Refactoring Dataset Inventory")
    print("--------------------------------")
    print(f"Metadata records: {len(records)}")
    print(f"Valid before/after pairs: {statuses['valid_pair']}")
    print(f"Missing or empty pairs: {statuses['missing_or_empty']}")

    print("\nRefactoring categories:")

    for category, count in sorted(categories.items()):
        print(f"  {category}: {count}")

    print(f"\nInventory saved to: {OUTPUT_FILE}")


if __name__ == "__main__":
    build_inventory()
